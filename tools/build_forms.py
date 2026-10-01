#!/usr/bin/env python3
"""Build the FieldComm Forms library: fillable PDFs + Excel logs.

    python3 tools/build_forms.py      # writes forms/*.pdf and forms/*.xlsx

Each form is defined below as plain data (header fields, tables, text boxes).
Add a form by adding a definition and a line to FORMS. Requires reportlab and
openpyxl.
"""
import os

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from reportlab.lib.colors import HexColor, white
from reportlab.lib.pagesizes import landscape, letter
from reportlab.pdfgen import canvas

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
OUT = os.path.join(ROOT, "forms")

NAVY = HexColor("#0D1B2A")
GOLD = HexColor("#C9A84C")
INK = HexColor("#1B2430")
GRID = HexColor("#9AA5B1")
SHADE = HexColor("#EEF1F4")
REV = "Rev 1.0"
FOOTER = ("FieldComm Forms · free to download, use and adapt · "
          "Template only — not legal advice; verify requirements with your contract, AHJ and attorney.")
M = 36  # page margin (0.5 in)


class Page:
    """One PDF page with a cursor that moves down as blocks are added."""

    def __init__(self, c, size, code, title, subtitle):
        self.c, (self.w, self.h) = c, size
        c.setPageSize(size)
        self.code, self.title = code, title
        self.y = self.h - M
        self._n = 0
        self._header(subtitle)

    def name(self, base):
        self._n += 1
        return f"{self.title}_{base}_{self._n}".replace(" ", "_")

    def _header(self, subtitle):
        c, w = self.c, self.w
        c.setFillColor(NAVY)
        c.rect(M, self.y - 46, w - 2 * M, 46, stroke=0, fill=1)
        c.setFillColor(GOLD)
        c.rect(M, self.y - 49, w - 2 * M, 3, stroke=0, fill=1)
        c.setFillColor(white)
        c.setFont("Helvetica-Bold", 17)
        c.drawString(M + 12, self.y - 22, self.title.upper())
        c.setFont("Helvetica", 8.5)
        c.drawString(M + 12, self.y - 36, subtitle)
        c.setFont("Helvetica-Bold", 10)
        c.drawRightString(w - M - 12, self.y - 20, "FIELDCOMM")
        c.setFillColor(GOLD)
        c.setFont("Helvetica", 8)
        c.drawRightString(w - M - 12, self.y - 33, f"{self.code} · {REV}")
        self.y -= 62
        # footer
        c.setFillColor(GRID)
        c.setFont("Helvetica", 6.5)
        c.drawString(M, M - 14, FOOTER)
        c.drawRightString(w - M, M - 4, f"{self.code} · {self.title}")

    def label(self, x, y, text):
        self.c.setFillColor(INK)
        self.c.setFont("Helvetica-Bold", 6.8)
        self.c.drawString(x, y, text.upper())

    def field(self, name, x, y, w, h=15, multiline=False, tooltip=""):
        flags = "multiline" if multiline else ""
        self.c.acroForm.textfield(
            name=name, tooltip=tooltip or name, x=x, y=y, width=w, height=h,
            borderWidth=0.6, borderColor=GRID, fillColor=white, textColor=INK,
            fontName="Helvetica", fontSize=0 if multiline else 9, fieldFlags=flags)

    def row(self, items):
        """items: list of (label, width_fraction)."""
        x, avail = M, self.w - 2 * M
        gap = 8
        widths = [f * (avail - gap * (len(items) - 1)) for _, f in items]
        for (lab, _), wd in zip(items, widths):
            self.label(x, self.y - 8, lab)
            self.field(self.name(lab), x, self.y - 27, wd, 16, tooltip=lab)
            x += wd + gap
        self.y -= 36

    def checks(self, label, options):
        self.label(M, self.y - 8, label)
        x = M
        for opt in options:
            self.c.acroForm.checkbox(name=self.name(f"{label}_{opt}"), tooltip=f"{label}: {opt}",
                                     x=x, y=self.y - 26, size=11, borderColor=GRID,
                                     fillColor=white, buttonStyle="check", borderWidth=0.6)
            self.c.setFont("Helvetica", 8.5)
            self.c.setFillColor(INK)
            self.c.drawString(x + 15, self.y - 23, opt)
            x += 15 + self.c.stringWidth(opt, "Helvetica", 8.5) + 16
        self.y -= 34

    def box(self, label, height):
        self.label(M, self.y - 8, label)
        self.field(self.name(label), M, self.y - 12 - height, self.w - 2 * M, height,
                   multiline=True, tooltip=label)
        self.y -= height + 20

    def table(self, title, cols, rows, row_h=17):
        """cols: list of (header, width_fraction)."""
        c = self.c
        avail = self.w - 2 * M
        widths = [f * avail for _, f in cols]
        if title:
            self.label(M, self.y - 8, title)
            self.y -= 12
        # header band (long headers wrap onto extra lines)
        from reportlab.lib.utils import simpleSplit
        fs = 6.4
        heads = [simpleSplit(h.upper(), "Helvetica-Bold", fs, wd - 5) or [""] for (h, _), wd in zip(cols, widths)]
        band = 8 + (fs + 1.4) * max(len(hl) for hl in heads)
        c.setFillColor(SHADE)
        c.rect(M, self.y - band, avail, band, stroke=0, fill=1)
        c.setFillColor(INK)
        c.setFont("Helvetica-Bold", fs)
        x = M
        for hl, wd in zip(heads, widths):
            for k, ln in enumerate(hl):
                c.drawString(x + 3, self.y - 4 - fs - k * (fs + 1.4), ln)
            x += wd
        self.y -= band
        for r in range(rows):
            x = M
            for (h, _), wd in zip(cols, widths):
                self.field(f"{self.title}_{title or 'row'}_{h}_{r + 1}".replace(" ", "_"), x, self.y - row_h, wd, row_h,
                           tooltip=f"{h} (row {r + 1})")
                x += wd
            self.y -= row_h
        self.y -= 12

    def signature(self, labels):
        x, avail = M, self.w - 2 * M
        wd = (avail - 12 * (len(labels) - 1)) / len(labels)
        for lab in labels:
            self.c.setStrokeColor(INK)
            self.c.setLineWidth(0.7)
            self.c.line(x, self.y - 22, x + wd, self.y - 22)
            self.label(x, self.y - 32, lab)
            x += wd + 12
        self.y -= 40

    def checklist(self, title, items, extra=("Date", "By")):
        """One line per item: tick box, label, then small fields (date / initials)."""
        c = self.c
        avail = self.w - 2 * M
        ex_w = 70
        if title:
            self.label(M, self.y - 8, title)
            self.y -= 14
        for i, item in enumerate(items, start=1):
            c.acroForm.checkbox(name=self.name(f"chk_{i}"), tooltip=item, x=M, y=self.y - 14,
                                size=11, borderColor=GRID, fillColor=white, buttonStyle="check",
                                borderWidth=0.6)
            c.setFillColor(INK)
            c.setFont("Helvetica", 8.6)
            c.drawString(M + 17, self.y - 11, item)
            x = M + avail - ex_w * len(extra) - 6 * (len(extra) - 1)
            for e in extra:
                c.setFont("Helvetica", 6)
                c.setFillColor(GRID)
                c.drawString(x + 2, self.y + 1, e.upper())
                self.field(self.name(f"{e}_{i}"), x, self.y - 15, ex_w, 13, tooltip=f"{item}: {e}")
                x += ex_w + 6
            c.setStrokeColor(SHADE)
            c.setLineWidth(0.6)
            c.line(M, self.y - 19, M + avail, self.y - 19)
            self.y -= 21
        self.y -= 8

    def para(self, text, size=8.4):
        """Fixed wording (e.g. an attestation statement), wrapped to the page width."""
        from reportlab.lib.utils import simpleSplit
        lines = simpleSplit(text, "Helvetica", size, self.w - 2 * M)
        self.c.setFillColor(INK)
        self.c.setFont("Helvetica", size)
        for ln in lines:
            self.c.drawString(M, self.y - size, ln)
            self.y -= size + 3
        self.y -= 8

    def note(self, text):
        self.c.setFillColor(INK)
        self.c.setFont("Helvetica-Oblique", 8)
        self.c.drawString(M, self.y - 9, text)
        self.y -= 16


# ---------------------------------------------------------------- forms

def daily_report(c):
    p = Page(c, letter, "FC-FORM-001", "Daily Report", "One page per day. Fills the gap between memory and the record.")
    p.row([("Project", .46), ("Report #", .14), ("Date", .18), ("Superintendent", .22)])
    p.row([("Weather AM", .25), ("Weather PM", .25), ("High / Low °F", .2), ("Precip. / Wind", .3)])
    p.checks("Work stopped for weather?", ["No", "Partial", "Yes — full day"])
    p.table("Crew on site", [("Company", .26), ("Trade", .16), ("Workers", .09), ("Hours", .08), ("Work performed / location", .41)], 8)
    p.table("Equipment on site", [("Equipment", .4), ("Owner / rental", .3), ("Hours used", .15), ("Idle?", .15)], 3)
    p.box("Deliveries received (supplier, material, ticket #, short/damaged?)", 40)
    p.box("Visitors & inspections (name, agency, result)", 34)
    p.box("Delays, issues, RFIs or changes raised today", 40)
    p.box("Safety: toolbox talk topic, incidents, near misses", 30)
    p.signature(["Superintendent signature", "Date"])
    p.c.showPage()


def rfi_request(c):
    p = Page(c, letter, "FC-FORM-002", "Request for Information", "Ask one clear question. Reference the sheet and spec. Propose an answer.")
    p.row([("Project", .5), ("RFI #", .15), ("Date sent", .17), ("Response needed by", .18)])
    p.row([("To (design team)", .5), ("From", .5)])
    p.row([("Subject", .6), ("Drawing / sheet ref.", .2), ("Spec section", .2)])
    p.box("Question", 120)
    p.box("Suggested solution (contractor's proposal)", 90)
    p.checks("Potential impact", ["Cost", "Schedule", "None expected", "Unknown"])
    p.box("Response (design team)", 140)
    p.row([("Responded by", .5), ("Date answered", .25), ("Distributed to", .25)])
    p.signature(["Requested by", "Answered by"])
    c.showPage()


def rfi_log(c):
    p = Page(c, landscape(letter), "FC-FORM-002", "RFI Log", "Every RFI, one line. Open items stay visible until they're closed.")
    p.row([("Project", .5), ("Log kept by", .3), ("Page", .2)])
    p.table(None, [("RFI #", .06), ("Date sent", .08), ("Subject", .25), ("Sheet / spec", .1), ("To", .11),
                   ("Needed by", .08), ("Answered", .08), ("Status", .08), ("Cost/Sched impact", .16)], 22, row_h=17)
    c.showPage()


def punch_list(c):
    p = Page(c, landscape(letter), "FC-FORM-003", "Punch List", "Walk it, write it, assign it, verify it. Nothing closes until someone checks it.")
    p.row([("Project", .4), ("Area / building / unit", .3), ("Walk date", .15), ("Walked by", .15)])
    p.table(None, [("#", .04), ("Location", .13), ("Trade", .1), ("Item / deficiency", .31), ("Assigned to", .12),
                   ("Due", .07), ("Done", .07), ("Verified by", .16)], 22, row_h=17)
    c.showPage()


def change_order_log(c):
    p = Page(c, landscape(letter), "FC-FORM-004", "Change Order Log", "Track every change from first mention to signed approval. Cost and days, both.")
    p.row([("Project", .45), ("Original contract $", .2), ("Original completion date", .2), ("Page", .15)])
    p.table(None, [("CO #", .06), ("Ref (RFI / PCO)", .09), ("Description", .29), ("Requested by", .11), ("Date", .08),
                   ("Cost +/-", .1), ("Days +/-", .07), ("Status", .1), ("Approved", .1)], 18, row_h=17)
    p.row([("Total approved $", .25), ("Total approved days", .25), ("Revised contract $", .25), ("Revised completion", .25)])
    c.showPage()


def change_order_request(c):
    p = Page(c, letter, "FC-FORM-004", "Change Order Request", "Describe the change, why it's needed, and what it costs in money and time.")
    p.row([("Project", .5), ("CO request #", .2), ("Date", .3)])
    p.row([("Requested by", .5), ("Related RFI / bulletin / directive", .5)])
    p.checks("Reason", ["Owner request", "Design change", "Unforeseen condition", "Code / AHJ", "Other"])
    p.box("Description of change", 100)
    p.table("Cost breakdown", [("Item", .44), ("Qty", .1), ("Unit", .1), ("Unit cost", .16), ("Amount", .2)], 7)
    p.row([("Subtotal $", .25), ("Overhead & profit $", .25), ("Total change $", .25), ("Schedule impact (days)", .25)])
    p.signature(["Contractor", "Owner / owner's rep", "Date approved"])
    c.showPage()


def toolbox_talk(c):
    p = Page(c, letter, "FC-FORM-005", "Toolbox Talk Sign-In", "Short talk, real topic, every worker signs. The sign-in is your proof it happened.")
    p.row([("Project", .45), ("Date", .18), ("Time", .12), ("Presented by", .25)])
    p.row([("Topic", .7), ("Language(s) used", .3)])
    p.box("Key points covered", 70)
    p.table("Attendance", [("#", .05), ("Print name", .32), ("Company", .25), ("Trade", .14), ("Signature", .24)], 18, row_h=18)
    p.signature(["Presenter signature", "Superintendent"])
    c.showPage()



def photo_log(c):
    p = Page(c, landscape(letter), "FC-FORM-006", "Photo Log", "Number every photo, say where it was taken and what it shows. A photo with no location is just a picture.")
    p.row([("Project", .4), ("Building / unit", .25), ("Photographer", .2), ("Page", .15)])
    p.table(None, [("Photo #", .07), ("Date", .08), ("Time", .06), ("Location (bldg / unit / room)", .2), ("Trade", .1),
                   ("View / direction", .1), ("What it shows", .27), ("Linked to (daily / RFI / punch #)", .12)], 22, row_h=17)
    c.showPage()


def look_ahead(c):
    p = Page(c, landscape(letter), "FC-FORM-007", "3-Week Look-Ahead", "What happens in the next 15 working days, who does it, and what could stop it. Mark the days with an X.")
    p.row([("Project", .4), ("Week 1 starts (Mon)", .2), ("Prepared by", .25), ("Date issued", .15)])
    cols = [("Activity", .2), ("Trade / sub", .1), ("Area", .08)] + \
           [(f"{'MTWTF'[i % 5]}{i // 5 + 1}", .028) for i in range(15)] + [("Constraints / needs", .2)]
    p.note("Day columns: M1–F1 = week 1, M2–F2 = week 2, M3–F3 = week 3.")
    p.table(None, cols, 20, row_h=17)
    c.showPage()


def receiving_log(c):
    p = Page(c, landscape(letter), "FC-FORM-008", "Delivery & Material Receiving Log", "Count it at the tailgate. Short, damaged or wrong material gets written down before the truck leaves.")
    p.row([("Project", .45), ("Log kept by", .3), ("Page", .25)])
    p.table(None, [("Date", .07), ("Supplier", .13), ("Ticket / PO #", .1), ("Material", .2), ("Qty ordered", .08),
                   ("Qty received", .08), ("Short / back-ordered", .1), ("Damaged?", .07), ("Received by", .09), ("Stored at", .08)],
            22, row_h=17)
    c.showPage()


def credential_check(c):
    p = Page(c, letter, "FC-FORM-009", "Credential & Scope Authorization", "Right credential, right qualification, still valid — checked before the work starts.")
    p.para("An unexpired card is not the same as authorization. Check the credential against what this specific scope "
           "requires (for welding: process, position, material and thickness range), confirm it is current for the whole "
           "duration of the work, and record who verified it. If it doesn't match, the work does not start.")
    p.row([("Project", .4), ("Scope / work item", .35), ("Date", .25)])
    p.row([("Location (bldg / level / area)", .5), ("Subcontractor", .5)])
    p.row([("Required credential type", .5), ("Qualification needed (process / position / material / thickness)", .5)])
    p.checks("Requirement comes from", ["Spec section", "Special inspection schedule", "AHJ / permit", "Contract", "Other"])
    p.table("Workers assigned to this scope", [("Name", .18), ("Credential type", .16), ("Cert #", .1), ("Issued by", .12),
                                               ("Expires", .09), ("Qualification matches scope? (Y/N)", .17), ("Verified by / how", .18)], 7)
    p.checks("Checked", ["Original or issuer lookup, not just a photo", "Recent-use / continuity confirmed", "Valid through scope completion"])
    p.checks("Decision", ["AUTHORIZED to proceed", "HOLD — mismatch", "HOLD — expired", "HOLD — missing"])
    p.row([("Scope expected to finish", .3), ("Earliest credential expiry on this scope", .35), ("Re-check date", .35)])
    p.box("Notes / corrective action", 46)
    p.signature(["Superintendent", "Verified by", "Date"])
    c.showPage()


def incident_report(c):
    p = Page(c, letter, "FC-FORM-010", "Incident & Near-Miss Report", "Write it the same day. Near misses count — they are the warning before the injury.")
    p.row([("Project", .4), ("Date", .18), ("Time", .14), ("Exact location", .28)])
    p.checks("Type", ["Near miss", "First aid", "Recordable injury", "Property damage", "Environmental"])
    p.table("People involved", [("Name", .28), ("Company", .24), ("Trade", .16), ("Injury / role", .32)], 3)
    p.box("What happened (sequence of events)", 70)
    p.box("Immediate cause (what directly caused it)", 34)
    p.box("Root cause (why the conditions existed)", 34)
    p.table("Witnesses", [("Name", .4), ("Company", .35), ("Phone", .25)], 3)
    p.table("Corrective actions", [("Action", .52), ("Owner", .2), ("Due", .14), ("Done", .14)], 3)
    p.note("Recordable injuries must also be entered on the official OSHA 300 / 301 forms (osha.gov).")
    p.signature(["Reported by", "Superintendent", "Date"])
    c.showPage()


def completion_attestation(c):
    p = Page(c, letter, "FC-FORM-011", "Work Completion Attestation", "Verify the work before the money moves. Each line of the pay request is backed by proof.")
    p.row([("Project", .4), ("Subcontractor", .35), ("Pay period", .25)])
    p.row([("Scope / schedule-of-values line", .6), ("Location", .4)])
    p.box("Work completed this period", 66)
    p.row([("% complete — previous", .25), ("% this period", .25), ("% complete to date", .25), ("$ requested this period", .25)])
    p.checklist("Verification", [
        "Superintendent inspected the work in person",
        "Photos attached (photo log #s)",
        "Required inspection passed (permit / inspection #)",
        "Punch items for this scope closed or listed below",
        "Credentials verified for credential-required scope (FC-FORM-009)",
        "Lien waiver received for prior payment",
    ], extra=("Date", "Initials"))
    p.box("Exceptions / open items", 40)
    p.para("The undersigned attests that the work described above has been performed in accordance with the contract "
           "documents to the percentage stated, and that the information on this form is true and complete to the best of "
           "their knowledge.")
    p.checks("Release for payment", ["Yes — release this line", "Partial", "HOLD"])
    p.signature(["Subcontractor", "Superintendent (verified)", "Date"])
    c.showPage()


def coi_lien_tracker(c):
    p = Page(c, landscape(letter), "FC-FORM-012", "COI & Lien Waiver Tracker", "No current insurance, no work. No waiver for the last payment, no next payment.")
    p.row([("Project", .45), ("Tracked by", .3), ("Page", .25)])
    p.table(None, [("Subcontractor", .14), ("GL exp.", .07), ("Auto exp.", .07), ("WC exp.", .07), ("Umbrella exp.", .07),
                   ("Add'l insured?", .07), ("Pay app # / period", .1), ("Amount paid", .09), ("Conditional waiver", .09),
                   ("Unconditional waiver", .09), ("Notes", .14)], 20, row_h=17)
    p.note("Use the lien waiver forms required by your state (in Florida, the statutory forms in Fla. Stat. 713.20).")
    c.showPage()


def closeout_checklist(c):
    p = Page(c, letter, "FC-FORM-013", "Closeout Checklist", "Everything the owner needs to take the building and release the last payment.")
    p.row([("Project", .5), ("Substantial completion date", .25), ("Checklist owner", .25)])
    p.checklist(None, [
        "Final inspections passed / permits closed",
        "Certificate of Occupancy (or Completion) issued",
        "Punch list complete and verified",
        "As-built drawings delivered",
        "O&M manuals delivered",
        "Warranties and guarantees collected",
        "Attic stock / spare materials turned over",
        "Owner training completed (systems, equipment)",
        "Keys, access cards and codes turned over",
        "Utilities transferred to owner",
        "Final lien waivers from all subs and suppliers",
        "Consent of surety (if bonded)",
        "Final change orders executed",
        "Final pay application submitted",
        "Retainage release requested",
        "Site cleaned, temporary facilities removed",
    ])
    p.box("Open items / notes", 50)
    p.signature(["Contractor", "Owner / owner's rep", "Date"])
    c.showPage()

FORMS = [
    ("FC-FORM-001_Daily-Report", [daily_report]),
    ("FC-FORM-002_RFI-Request-and-Log", [rfi_request, rfi_log]),
    ("FC-FORM-003_Punch-List", [punch_list]),
    ("FC-FORM-004_Change-Order-Request-and-Log", [change_order_request, change_order_log]),
    ("FC-FORM-005_Toolbox-Talk-Sign-In", [toolbox_talk]),
    ("FC-FORM-006_Photo-Log", [photo_log]),
    ("FC-FORM-007_3-Week-Look-Ahead", [look_ahead]),
    ("FC-FORM-008_Receiving-Log", [receiving_log]),
    ("FC-FORM-009_Credential-Scope-Authorization", [credential_check]),
    ("FC-FORM-010_Incident-Near-Miss-Report", [incident_report]),
    ("FC-FORM-011_Work-Completion-Attestation", [completion_attestation]),
    ("FC-FORM-012_COI-Lien-Waiver-Tracker", [coi_lien_tracker]),
    ("FC-FORM-013_Closeout-Checklist", [closeout_checklist]),
]

# Excel versions of the logs (same columns as the PDFs)
SHEETS = [
    ("FC-FORM-002_RFI-Log", "RFI Log",
     ["RFI #", "Date sent", "Subject", "Sheet / spec", "To", "Needed by", "Answered", "Status", "Cost/Sched impact"],
     [8, 12, 40, 14, 18, 12, 12, 12, 24], 60),
    ("FC-FORM-003_Punch-List", "Punch List",
     ["#", "Location", "Trade", "Item / deficiency", "Assigned to", "Due", "Done", "Verified by"],
     [6, 18, 14, 48, 18, 11, 11, 18], 100),
    ("FC-FORM-004_Change-Order-Log", "Change Order Log",
     ["CO #", "Ref (RFI / PCO)", "Description", "Requested by", "Date", "Cost +/-", "Days +/-", "Status", "Approved"],
     [8, 14, 44, 16, 12, 14, 10, 12, 12], 60),
    ("FC-FORM-006_Photo-Log", "Photo Log",
     ["Photo #", "Date", "Time", "Location (bldg / unit / room)", "Trade", "View / direction", "What it shows", "Linked to"],
     [9, 11, 8, 30, 14, 14, 44, 18], 200),
    ("FC-FORM-007_3-Week-Look-Ahead", "3-Week Look-Ahead",
     ["Activity", "Trade / sub", "Area"] + [f"{'MTWTF'[i % 5]}{i // 5 + 1}" for i in range(15)] + ["Constraints / needs"],
     [34, 16, 12] + [4] * 15 + [34], 40),
    ("FC-FORM-008_Receiving-Log", "Receiving Log",
     ["Date", "Supplier", "Ticket / PO #", "Material", "Qty ordered", "Qty received", "Short / back-ordered", "Damaged?",
      "Received by", "Stored at"],
     [11, 20, 14, 32, 11, 11, 16, 10, 14, 14], 150),
    ("FC-FORM-012_COI-Lien-Waiver-Tracker", "COI and Lien Waivers",
     ["Subcontractor", "GL exp.", "Auto exp.", "WC exp.", "Umbrella exp.", "Add'l insured?", "Pay app # / period",
      "Amount paid", "Conditional waiver", "Unconditional waiver", "Notes"],
     [26, 11, 11, 11, 12, 12, 16, 13, 14, 14, 30], 80),
]


def build_pdf(fname, pages):
    path = os.path.join(OUT, fname + ".pdf")
    c = canvas.Canvas(path, pagesize=letter)
    c.setTitle(fname.replace("_", " ").replace("-", " "))
    c.setAuthor("FieldComm")
    c.setSubject("Free construction field form")
    for page in pages:
        page(c)
    c.save()
    return path


def build_xlsx(fname, title, headers, widths, rows):
    wb = Workbook()
    ws = wb.active
    ws.title = title[:31]
    navy = PatternFill("solid", fgColor="0D1B2A")
    thin = Side(style="thin", color="9AA5B1")
    ws["A1"] = f"FIELDCOMM · {title.upper()}"
    ws["A1"].font = Font(bold=True, size=14, color="FFFFFF")
    ws["A2"] = "Project:"
    ws["A2"].font = Font(bold=True)
    for col in range(1, len(headers) + 1):
        ws.cell(row=1, column=col).fill = navy
    for i, (h, w) in enumerate(zip(headers, widths), start=1):
        cell = ws.cell(row=4, column=i, value=h)
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1E3A5F")
        cell.alignment = Alignment(vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = w
    for r in range(5, 5 + rows):
        for col in range(1, len(headers) + 1):
            ws.cell(row=r, column=col).border = Border(top=thin, bottom=thin, left=thin, right=thin)
    ws.freeze_panes = "A5"
    ws.auto_filter.ref = f"A4:{get_column_letter(len(headers))}{4 + rows}"
    ws.cell(row=6 + rows, column=1, value=FOOTER).font = Font(italic=True, size=8, color="667085")
    path = os.path.join(OUT, fname + ".xlsx")
    wb.save(path)
    return path


def main():
    os.makedirs(OUT, exist_ok=True)
    for fname, pages in FORMS:
        print(build_pdf(fname, pages))
    for args in SHEETS:
        print(build_xlsx(*args))


if __name__ == "__main__":
    main()
